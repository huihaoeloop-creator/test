#!/usr/bin/env python3
"""
Dump a PowerPoint template's structure: slide size, masters, layouts, and every
placeholder's index, type, name and geometry.

The placeholder idx values are the contract between the template and any script
that fills it — python-pptx addresses placeholders by idx, not by position, so
this listing is what a filling script is written against.

Real-world templates carry surprises (custom placeholder subtypes, layouts with
broken relationships, geometry inherited from the master), so every layout and
placeholder is read defensively: a failure is recorded in place and the dump
keeps going rather than aborting the whole run.

Usage:
    pip install python-pptx
    python -X utf8 dump_template.py 模板.potx            # 人类可读
    python -X utf8 dump_template.py 模板.potx --json     # 机器可读
"""
import argparse
import io
import json
import platform
import re
import sys
import traceback
import zipfile

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("请先安装依赖：pip install python-pptx")


# python-pptx accepts only the "presentation" main content type. A .potx (and
# .ppsx / .potm) is byte-for-byte the same package with a different type string,
# so rewrite that one attribute in memory and hand python-pptx the result. The
# file on disk is never modified.
MAIN_PART_CT = re.compile(
    r'ContentType="[^"]*(?:presentationml|ms-powerpoint)[^"]*\.main\+xml"'
)
PRESENTATION_CT = (
    'ContentType="application/vnd.openxmlformats-officedocument'
    '.presentationml.presentation.main+xml"'
)


def open_presentation(path):
    """Open a .pptx directly, or a template package via an in-memory rewrite.

    Returns (presentation, was_rewritten).
    """
    try:
        return Presentation(path), False
    except ValueError as exc:
        if "not a PowerPoint file" not in str(exc):
            raise

    buf = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = MAIN_PART_CT.sub(PRESENTATION_CT, data.decode("utf-8")).encode("utf-8")
            dst.writestr(item, data)
    buf.seek(0)
    return Presentation(buf), True


def inches(value):
    """Placeholders may inherit geometry from the master, leaving these unset."""
    try:
        return round(Emu(value).inches, 2) if value is not None else None
    except Exception:
        return None


def safe(fn, fallback=None):
    try:
        return fn()
    except Exception as exc:
        return f"<读取失败: {type(exc).__name__}: {exc}>" if fallback is None else fallback


def describe(shape):
    fmt = shape.placeholder_format
    # str(type) looks like "BODY (2)"; keep the readable half.
    kind = safe(lambda: str(fmt.type).split(" (")[0] if fmt.type is not None else "UNKNOWN")
    return {
        "idx": safe(lambda: fmt.idx),
        "type": kind,
        "name": safe(lambda: shape.name),
        "left_in": inches(safe(lambda: shape.left, fallback=None)),
        "top_in": inches(safe(lambda: shape.top, fallback=None)),
        "width_in": inches(safe(lambda: shape.width, fallback=None)),
        "height_in": inches(safe(lambda: shape.height, fallback=None)),
    }


def describe_layout(index, layout):
    info = {"index": index, "name": safe(lambda: layout.name), "placeholders": [], "error": None}
    try:
        for ph in layout.placeholders:
            try:
                info["placeholders"].append(describe(ph))
            except Exception as exc:
                info["placeholders"].append({"error": f"{type(exc).__name__}: {exc}"})
    except Exception as exc:
        info["error"] = f"{type(exc).__name__}: {exc}"
    return info


def collect(path):
    prs, rewritten = open_presentation(path)
    report = {
        "file": path,
        "python": platform.python_version(),
        "opened_as_template": rewritten,
        "slide_width_in": inches(prs.slide_width),
        "slide_height_in": inches(prs.slide_height),
        "masters": [],
        "existing_slides": [],
    }

    for m_i, master in enumerate(prs.slide_masters):
        m_info = {"index": m_i, "name": safe(lambda: master.name), "layouts": []}
        try:
            for l_i, layout in enumerate(master.slide_layouts):
                m_info["layouts"].append(describe_layout(l_i, layout))
        except Exception as exc:
            m_info["error"] = f"{type(exc).__name__}: {exc}"
        report["masters"].append(m_info)

    # A .potx normally carries no slides; a .pptx used as a template does.
    try:
        report["existing_slides"] = [
            {"index": i, "layout": safe(lambda: s.slide_layout.name)}
            for i, s in enumerate(prs.slides)
        ]
    except Exception as exc:
        report["existing_slides"] = [{"error": f"{type(exc).__name__}: {exc}"}]

    return report


def render(report):
    w, h = report["slide_width_in"], report["slide_height_in"]
    shape = "?"
    if isinstance(w, float) and isinstance(h, float) and h:
        ratio = w / h
        shape = "16:9" if abs(ratio - 16 / 9) < 0.05 else "4:3" if abs(ratio - 4 / 3) < 0.05 else f"{ratio:.2f}:1"

    print(f"文件: {report['file']}")
    print(f"Python: {report['python']}")
    if report.get("opened_as_template"):
        print("类型: 模板包（.potx/.ppsx），已在内存中转换后读取，原文件未改动")
    print(f"页面: {w} x {h} in  ({shape})")

    for master in report["masters"]:
        print(f"\n母版 {master['index']}: {master['name']}")
        if master.get("error"):
            print(f"  !! 读取版式失败: {master['error']}")
        for layout in master["layouts"]:
            print(f"\n  版式 [{layout['index']}] {layout['name']}")
            if layout.get("error"):
                print(f"      !! {layout['error']}")
            if not layout["placeholders"]:
                print("      (无占位符)")
            for ph in layout["placeholders"]:
                if "error" in ph:
                    print(f"      !! {ph['error']}")
                    continue
                pos = (
                    f"({ph['left_in']}, {ph['top_in']}) {ph['width_in']}x{ph['height_in']} in"
                    if ph["left_in"] is not None
                    else "(几何继承自母版)"
                )
                print(f"      idx={str(ph['idx']):<4} {str(ph['type']):<16} {ph['name']!r:<34} {pos}")

    if report["existing_slides"]:
        print(f"\n模板内已有 {len(report['existing_slides'])} 页幻灯片:")
        for s in report["existing_slides"]:
            if "error" in s:
                print(f"  !! {s['error']}")
            else:
                print(f"  第 {s['index'] + 1} 页 -> 版式 {s['layout']!r}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("template", help=".potx 或 .pptx 路径")
    parser.add_argument("--json", action="store_true", help="输出 JSON")
    args = parser.parse_args()

    try:
        data = collect(args.template)
    except Exception:
        print("解析模板失败，完整报错如下，请把这一整段发回：\n", file=sys.stderr)
        traceback.print_exc()
        sys.exit(1)

    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2, default=str))
    else:
        render(data)
